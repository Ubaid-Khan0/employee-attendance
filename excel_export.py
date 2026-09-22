import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from database import (
    get_employee_attendance,
    get_all_attendance,
    get_all_employees
)


EXPORT_DIR = "employee_reports"


def create_employee_excel(employee_name):
    os.makedirs(EXPORT_DIR, exist_ok=True)

    records = get_employee_attendance(
        employee_name
    )

    filename = (
        employee_name.replace(" ", "_")
        + ".xlsx"
    )

    filepath = os.path.join(
        EXPORT_DIR,
        filename
    )

    workbook = Workbook()
    sheet = workbook.active

    sheet.title = "Attendance"

    headers = [
        "Employee Name",
        "Date",
        "Check-in Time",
        "Check-out Time",
        "Status",
        "Late Minutes"
    ]

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(
            row=1,
            column=column,
            value=header
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center"
        )

    for row_number, record in enumerate(
        records,
        start=2
    ):
        for column, value in enumerate(
            record,
            start=1
        ):
            sheet.cell(
                row=row_number,
                column=column,
                value=value
            )

    for column in sheet.columns:
        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[
            column_letter
        ].width = max_length + 3

    workbook.save(filepath)

    return filepath


def create_company_excel():
    os.makedirs(EXPORT_DIR, exist_ok=True)

    filepath = os.path.join(
        EXPORT_DIR,
        "All_Employees_Attendance.xlsx"
    )

    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    employees = get_all_employees()

    # Individual employee sheets
    for employee in employees:

        employee_name = employee[0]

        sheet_name = employee_name[:31]

        sheet = workbook.create_sheet(
            title=sheet_name
        )

        headers = [
            "Employee Name",
            "Date",
            "Check-in Time",
            "Check-out Time",
            "Status",
            "Late Minutes"
        ]

        for column, header in enumerate(
            headers,
            start=1
        ):
            cell = sheet.cell(
                row=1,
                column=column,
                value=header
            )

            cell.font = Font(bold=True)

        records = get_employee_attendance(
            employee_name
        )

        for row_number, record in enumerate(
            records,
            start=2
        ):
            for column, value in enumerate(
                record,
                start=1
            ):
                sheet.cell(
                    row=row_number,
                    column=column,
                    value=value
                )

    # All employees sheet
    all_sheet = workbook.create_sheet(
        title="All Employees"
    )

    headers = [
        "Employee Name",
        "Date",
        "Check-in Time",
        "Check-out Time",
        "Status",
        "Late Minutes"
    ]

    for column, header in enumerate(
        headers,
        start=1
    ):
        cell = all_sheet.cell(
            row=1,
            column=column,
            value=header
        )

        cell.font = Font(bold=True)

    records = get_all_attendance()

    for row_number, record in enumerate(
        records,
        start=2
    ):
        for column, value in enumerate(
            record,
            start=1
        ):
            all_sheet.cell(
                row=row_number,
                column=column,
                value=value
            )

    workbook.save(filepath)

    return filepath